from docx import Document
from copy import deepcopy
from pathlib import Path

SRC = Path(r"D:\UMG\Purificadora\Capitulo_4_Analisis_Sistema_Purificadora.docx")
OUT = Path(r"D:\UMG\Purificadora\Capitulo_4_Analisis_Sistema_Purificadora_Humanizado.docx")

def replace_text(p, text):
    props = deepcopy(p.runs[0]._r.rPr) if p.runs else None
    p._element.clear_content()
    run = p.add_run(text)
    if props is not None:
        run._r.get_or_add_rPr().append(props)

doc = Document(SRC)

revisions = {
2: "Este proyecto propone un sistema para ordenar y controlar el trabajo diario de una distribuidora de agua pura. Nace de una necesidad concreta: cuando clientes, rutas, inventario, cargas, ventas y cobros se llevan de forma manual, es fácil que la información quede dispersa, se repita o no esté disponible cuando se necesita.",
3: "La solución se plantea como una aplicación empresarial con un backend en Java y Spring Boot, una interfaz web desarrollada en React y una base de datos PostgreSQL. La parte móvil funcionará como una aplicación web progresiva, de modo que los vendedores puedan continuar trabajando desde su teléfono aun si la conexión a internet falla por momentos.",
4: "El objetivo es que la purificadora pueda seguir sus operaciones con mayor claridad y respaldo. El sistema reunirá la configuración de la empresa, usuarios, productos, precios, clientes, rutas, cargas, ventas, pagos, créditos, devoluciones, mermas, liquidaciones, reportes y auditorías. No reemplaza las decisiones de la administración; más bien, les brinda información confiable para tomarlas.",
8: "Para esta primera versión no se incluyen varias empresas en una sola instalación, facturación electrónica integrada, rastreo permanente de vendedores, cobros con tarjeta ni decisiones administrativas automáticas. De esa forma, el esfuerzo se concentra en lo esencial: llevar un control confiable de la distribución de agua pura.",
10: "Los usuarios se definieron según la función que desempeñan dentro de la purificadora y el nivel de responsabilidad que tienen. El sistema trabajará con roles y permisos por recurso; es decir, cada persona podrá ver y realizar únicamente lo que necesita para cumplir con su trabajo.",
14: "Antes de poner el sistema en marcha, la administración deberá confirmar quiénes serán los usuarios autorizados, qué rutas atenderán, qué vehículos usarán y qué permisos tendrá cada persona. Los controles del sistema ayudan a aplicar esas decisiones de forma uniforme, aunque la supervisión operativa seguirá siendo responsabilidad de la empresa.",
16: "El modelado del negocio ayuda a entender cómo se trabaja actualmente y cómo cambiará la operación con el sistema. Se revisan las personas involucradas, los movimientos de producto y dinero, los registros que se generan y los puntos donde hace falta mayor control. Esta información sirve como base para definir requisitos, reglas de negocio, casos de uso y pruebas.",
18: "Actualmente, las cargas, ventas, devoluciones y cobros suelen anotarse a mano o reunirse después en hojas de cálculo. Al terminar la ruta, la administración debe revisar esos registros para conocer las existencias, el efectivo recibido y las posibles diferencias. Por eso, el resultado depende mucho de que cada anotación sea completa y de que alguien consolide la información correctamente.",
21: "La figura resume el recorrido habitual: se prepara la carga, el vendedor sale a ruta, realiza ventas y cobros, y al final se revisan los registros para hacer la liquidación.",
23: "En este proceso los principales problemas son el tiempo que toma reunir la información, la dificultad para conocer el inventario real de cada ruta y el riesgo de perder, duplicar o confundir comprobantes. También se vuelve más complicado consultar el historial de un cliente, producto, vendedor o fecha cuando la información está repartida en distintos registros.",
25: "Con el sistema propuesto, la carga, las ventas y la liquidación quedan conectadas en un mismo flujo. La aplicación valida permisos, existencias, precios, crédito y reglas de negocio desde el servidor. Si no hay internet, el vendedor puede conservar las operaciones autorizadas en su dispositivo y sincronizarlas después sin duplicar registros.",
28: "La figura muestra el recorrido esperado desde que bodega prepara la carga hasta que la administración revisa y cierra la liquidación. Un aspecto importante es que el inventario físico y el dinero se revisan por separado, por lo que una merma no disminuye el efectivo que el vendedor debe entregar.",
34: "Las reglas de negocio son acuerdos que el sistema debe respetar en cada operación. Se aplican principalmente en el servidor para que no dependan solo de lo que se muestra en pantalla. Además, las acciones importantes quedan registradas en auditoría y, cuando se requiere corregir algo, se usan anulaciones o movimientos compensatorios con autorización.",
39: "Para desarrollar el sistema se eligió una metodología ágil basada en Scrum y ajustada a las condiciones de un proyecto académico. Este enfoque permite avanzar por entregas pequeñas, revisar cada resultado a tiempo y dar prioridad a las funciones que más ayudan al control de la distribución.",
40: "La metodología se adapta al tamaño real del proyecto. El estudiante desarrollador realiza el trabajo técnico; la administración de la purificadora valida las reglas de operación y los resultados esperados; y el asesor acompaña la revisión académica y técnica. El product backlog se organiza a partir de los requisitos, las reglas de negocio y los riesgos identificados.",
45: "Las fases se trabajarán de forma iterativa. Cada incremento debe terminar con una función integrada, pruebas realizadas, evidencia técnica y documentación actualizada. La tabla presenta el orden de trabajo sin fijar fechas que todavía no hayan sido aprobadas en el cronograma del proyecto.",
59: "Dar prioridad a las funciones que más valor aportan a la operación: configuración, productos, clientes, rutas, inventario, carga, venta y liquidación.",
60: "Conservar la relación entre objetivos, requisitos, reglas de negocio, pruebas y evidencias de ejecución.",
61: "Revisar desde el inicio las reglas de operación con la administración, especialmente las relacionadas con precios, crédito, cargas, devoluciones, mermas y cierre de rutas.",
62: "Cuidar la información comercial mediante autenticación, permisos, autorización por recurso, auditoría y respaldos que puedan verificarse.",
63: "Documentar los cambios funcionales y técnicos para facilitar la revisión académica, la instalación y el mantenimiento posterior.",
65: "Los artefactos reúnen la información necesaria para construir y comprobar el sistema. Funcionan como una guía durante el desarrollo y permiten verificar que las funciones implementadas respondan a las necesidades reales de la purificadora. Si el alcance cambia, las prioridades podrán ajustarse siempre que el cambio quede documentado y aprobado.",
74: "Los valores de rendimiento, capacidad, conectividad y volumen de operaciones deberán revisarse con la purificadora y el asesor antes de establecerse como criterios definitivos. Así se evita fijar metas sin conocer las condiciones reales en las que se probará y utilizará el sistema.",
76: "El product backlog organiza las necesidades del sistema según su prioridad. Cada historia se escribe desde el punto de vista de la persona que recibirá el beneficio y se acompaña de un criterio de aceptación que permite comprobar si funciona como se espera. Después, esas historias pueden dividirse en tareas técnicas sin perder su relación con el requisito original.",
81: "Este análisis deja definida la base funcional, operativa y metodológica del sistema para la purificadora de agua. La propuesta reúne los procesos de distribución, permite dar seguimiento al inventario y al dinero, y facilita que la administración consulte información sobre clientes, rutas, ventas y liquidaciones.",
82: "Los procesos, requisitos, reglas de negocio y backlog descritos aquí servirán de apoyo para el Capítulo 5. Antes de finalizar la implementación y medir resultados, la purificadora deberá validar usuarios, políticas de precio y crédito, rutas, inventario inicial y el periodo de pruebas."
}
for index, text in revisions.items():
    replace_text(doc.paragraphs[index], text)

doc.save(OUT)
print(OUT)
