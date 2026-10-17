import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_manual():
    doc = docx.Document()

    # Configurar márgenes
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Colores de marca
    COLOR_PRIMARY = RGBColor(15, 41, 74)     # Azul marino corporativo #0F294A
    COLOR_SECONDARY = RGBColor(27, 108, 168) # Azul océano #1B6CA8
    COLOR_TEXT = RGBColor(40, 44, 52)        # Gris carbón oscuro
    COLOR_MUTED = RGBColor(100, 110, 125)    # Gris secundario

    # Configurar estilo normal
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = COLOR_TEXT

    # ------------------ PORTADA ELEGANTE ------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(80)
    title_p.paragraph_format.space_after = Pt(8)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = title_p.add_run("SISTEMA DE GESTIÓN Y DISTRIBUCIÓN")
    run_sub.font.size = Pt(12)
    run_sub.font.bold = True
    run_sub.font.color.rgb = COLOR_SECONDARY

    main_title_p = doc.add_paragraph()
    main_title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    main_title_p.paragraph_format.space_after = Pt(20)
    run_title = main_title_p.add_run("MANUAL DE USUARIO\nPOR ROLES OPERATIVOS")
    run_title.font.size = Pt(26)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    desc_p = doc.add_paragraph()
    desc_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    desc_p.paragraph_format.space_after = Pt(120)
    run_desc = desc_p.add_run("Guía paso a paso para los roles de:\n• ROL VENDEDOR (Ruta y Distribución Móvil)\n• ROL ADMINISTRADOR (Gestión Central y Supervisión)")
    run_desc.font.size = Pt(13)
    run_desc.font.italic = True
    run_desc.font.color.rgb = COLOR_MUTED

    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_meta = meta_p.add_run("Purificadora de Agua Pura\nVersión 2.0 • Edición Oficial 2026")
    run_meta.font.size = Pt(10)
    run_meta.font.bold = True
    run_meta.font.color.rgb = COLOR_SECONDARY

    doc.add_page_break()

    # Helper de títulos
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = COLOR_SECONDARY
        return p

    def add_p(text, bold_prefix=None, space_after=4):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.bold = True
            r_bold.font.color.rgb = COLOR_PRIMARY
        p.add_run(text)
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_b = p.add_run(bold_prefix)
            r_b.font.bold = True
            r_b.font.color.rgb = COLOR_PRIMARY
        p.add_run(text)
        return p

    def add_callout(text, title="NOTA IMPORTANTE:"):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_background(cell, "EBF3FA")
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r_title = p.add_run(title + " ")
        r_title.font.bold = True
        r_title.font.color.rgb = COLOR_SECONDARY
        r_text = p.add_run(text)
        r_text.font.size = Pt(10.5)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # =========================================================================
    # PARTE 1: ROL VENDEDOR
    # =========================================================================
    sec1 = doc.add_paragraph()
    sec1.paragraph_format.space_before = Pt(10)
    sec1.paragraph_format.space_after = Pt(8)
    r1 = sec1.add_run("SECCIÓN I: MANUAL OPERATIVO DEL VENDEDOR")
    r1.font.size = Pt(20)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_PRIMARY

    add_p("Esta sección detalla de forma exclusiva las actividades diarias, responsabilidades y uso de cada pantalla asignada al rol VENDEDOR, tanto desde dispositivos móviles (teléfono / tablet) como desde una computadora de escritorio.")

    add_h1("1. Flujo General de Trabajo Diario del Vendedor")
    add_p("El trabajo diario del vendedor sigue un ciclo continuo de 4 fases obligatorias:")
    
    # Tabla resumen de flujo
    table_v = doc.add_table(rows=5, cols=3)
    table_v.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Fase", "Módulo en el Sistema", "Acción Principal"]
    widths = [Inches(1.2), Inches(2.2), Inches(3.1)]

    for i, h in enumerate(headers):
        cell = table_v.cell(0, i)
        set_cell_background(cell, "0F294A")
        set_cell_margins(cell, 100, 100, 120, 120)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    data_v = [
        ("1. Inicio de Ruta", "Mi Carga (/loads)", "Verificar que el producto físico en el camión coincida con la carga registrada."),
        ("2. En Camino", "Mi Recorrido (/itinerary)", "Consultar la lista de clientes del día en el mapa ordenado y dirigirse a cada punto."),
        ("3. En el Cliente", "Ventas (/sales)", "Registrar venta, cobrar, registrar garrafones vacíos recibidos o prestados."),
        ("4. Fin de Día", "Liquidaciones (/settlements)", "Cuadrar el dinero total cobrado y el producto/envases sobrantes para entrega a bodega.")
    ]

    for row_idx, data_row in enumerate(data_v, start=1):
        bg = "F4F7FA" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data_row):
            cell = table_v.cell(row_idx, col_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, 80, 80, 120, 120)
            p = cell.paragraphs[0]
            p.add_run(text)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_h1("2. Paso a Paso en los Módulos del Vendedor")

    add_h2("2.1. Panel Operativo (/ - Inicio)")
    add_p("Es la pantalla de control principal del vendedor:")
    add_bullet("Muestra el resumen de Ventas de Hoy (monto en Quetzales facturado en la jornada).", "• Ventas del Día: ")
    add_bullet("Muestra la cantidad exacta de dinero que el vendedor tiene en su posesión física.", "• Efectivo en Mano: ")
    add_bullet("Muestra en tiempo real cuánto producto fue cargado, cuánto se ha vendido y cuánto queda disponible en el camión.", "• Stock en Camión: ")
    add_bullet("Acceso directo con un clic a 'Mi Recorrido' para ver el mapa de ruta y la lista de clientes.", "• Botón Ver Mi Ruta: ")

    add_h2("2.2. Verificación de Carga Asignada (/loads - 'Mi Carga')")
    add_p("Antes de salir a vender, el vendedor debe confirmar su inventario:")
    add_bullet("Ingresa al menú y selecciona 'Mi Carga'.")
    add_bullet("En la lista de cargas registradas aparecerá únicamente tu carga activa sin liquidar.")
    add_bullet("Comprueba las cantidades de Garrafones de 20L, botellas o galones asignados a tu camión.")
    add_bullet("Si la carga física coincide con la pantalla, estás listo para salir a tu ruta.")

    add_h2("2.3. Módulo 'Mi Recorrido' y Gestión de Clientes (/itinerary)")
    add_p("Este módulo guía la visita de todos los clientes asignados en tu ruta:")
    add_bullet("Los clientes aparecen numerados en el orden sugerido de entrega para minimizar el tiempo de traslado.", "1. Orden de Entrega: ")
    add_bullet("Cada cliente muestra su nombre, dirección y saldo previo si tuviera cuenta pendiente.", "2. Información Visible: ")
    add_bullet("Al presionar 'Ver Mapa', se desplegará el mapa geolocalizado con pines interactivos:", "3. Estados Visuales en Mapa y Lista: ")
    add_bullet("Cliente atendido con venta completada exitosamente hoy.", "   - Color Verde: ")
    add_bullet("Cliente pendiente de visita.", "   - Color Gris / Pendiente: ")
    add_bullet("Al pulsar sobre un cliente o presionar el botón 'Vender', el sistema te traslada directamente a la pantalla de Ventas con los datos del cliente ya cargados automáticamente.", "4. Atajo Rápido a Ventas: ")
    add_bullet("Si el cliente está cerrado o no compró, pulsa 'Visita sin Venta' e indica el motivo para que la supervisión sepa que pasaste por el lugar.", "5. Registro de No Venta: ")

    add_h2("2.4. Registro de Ventas (/sales)")
    add_p("Es el módulo más importante durante la atención en calle:")
    add_bullet("Si entraste desde 'Mi Recorrido', el cliente ya estará listo. Si no, selecciónalo del buscador.", "Paso 1 (Cliente): ")
    add_bullet("Selecciona la presentación (ej. Garrafón 20L) e indica la cantidad de unidades vendidas. El total se calculará de inmediato.", "Paso 2 (Productos): ")
    add_bullet("Ingresa cuántos garrafones vacíos te entrega el cliente. Si le dejas garrafones adicionales en calidad de préstamo, regístralo aquí.", "Paso 3 (Control de Garrafones): ")
    add_bullet("Elige Efectivo (indica con cuánto paga y el sistema calcula el cambio), Transferencia (ingresa el código de boleta) o Crédito (si está autorizado).", "Paso 4 (Método de Pago): ")
    add_bullet("Presiona 'Confirmar Venta'. El inventario de tu camión se descuenta de forma instantánea y el dinero se suma a tu saldo en mano.", "Paso 5 (Confirmación): ")

    add_h2("2.5. Mermas y Devoluciones (/wastes y /returns)")
    add_p("Manejo de eventualidades en el camino:")
    add_bullet("Si se quiebra o fisura un garrafón por bache o defecto, regístralo en 'Mermas' indicando la causa y cantidad. Esto justifica que no regrese lleno ni con cobro.", "• Mermas: ")
    add_bullet("Si un cliente rechaza un producto antes de abrirlo, se registra como devolución para que regrese a tu stock.", "• Devoluciones: ")

    add_h2("2.6. Liquidación y Cierre de Jornada (/settlements)")
    add_p("Al terminar el día y regresar a la planta de agua:")
    add_bullet("Ingresa a 'Liquidaciones'.")
    add_bullet("El sistema presentará el resumen: Total Cargado, Total Vendido, Mermas y Saldo Restante.")
    add_bullet("Verifica el dinero total en efectivo que debes entregar a caja central.")
    add_bullet("Entrega los garrafones sobrantes y vacíos a bodega y envía tu liquidación para firma y aprobación de la administración.")

    add_callout("El vendedor NUNCA debe dejar una carga sin liquidar al final del día. La liquidación diaria garantiza que tu stock del día siguiente empiece limpio y sin inconsistencias.", "REGLA DE ORO DEL VENDEDOR:")

    doc.add_page_break()

    # =========================================================================
    # PARTE 2: ROL ADMINISTRADOR
    # =========================================================================
    sec2 = doc.add_paragraph()
    sec2.paragraph_format.space_before = Pt(10)
    sec2.paragraph_format.space_after = Pt(8)
    r2 = sec2.add_run("SECCIÓN II: MANUAL OPERATIVO DEL ADMINISTRADOR")
    r2.font.size = Pt(20)
    r2.font.bold = True
    r2.font.color.rgb = COLOR_PRIMARY

    add_p("Esta sección detalla de forma exhaustiva las funciones de control, parametrización, supervisión financiera, inventarios y auditoría que corresponden exclusivamente al rol de ADMINISTRADOR.")

    add_h1("1. Panel Operativo Gerencial (/ - Inicio)")
    add_p("El administrador cuenta con una vista global de la empresa en tiempo real:")
    add_bullet("Ingresos totales facturados hoy (efectivo, transferencias, cuentas por cobrar).", "• Facturación Global: ")
    add_bullet("Monitoreo de qué camiones están en calle, cuáles han liquidado y cuáles tienen cargas pendientes.", "• Estado de Rutas: ")
    add_bullet("Total de agua purificada disponible para despacho en planta.", "• Nivel de Bodega Central: ")
    add_bullet("Avisos de transferencias bancarias por validar y solicitudes de anulación pendientes.", "• Alertas del Día: ")

    add_h1("2. Catálogo y Planificación del Negocio")

    add_h2("2.1. Presentaciones y Productos (/presentations y /products)")
    add_bullet("Registrar las capacidades de envases (Garrafón 20L, Galón 3.78L, Botella 500ml, etc.).", "1. Presentaciones: ")
    add_bullet("Indicar si el envase es retornable (requiere intercambio de vacío) o desechable.", "2. Tipo de Envase: ")
    add_bullet("Creación de artículos comercializables vinculados a sus presentaciones y códigos únicos.", "3. Productos: ")

    add_h2("2.2. Gestión de Clientes (/customers)")
    add_p("Control del padrón de clientes y puntos de entrega:")
    add_bullet("Nombre comercial, nombre del encargado, teléfonos y dirección.", "• Datos de Contacto: ")
    add_bullet("Al dar de alta o editar un cliente, se asigna su coordenada exacta en el mapa satelital. Esto garantiza que la ruta del vendedor funcione con geolocalización.", "• Ubicación GPS: ")
    add_bullet("Se establece qué días pasa el camión por su negocio/casa (ej. Lunes-Miércoles-Viernes).", "• Días de Visita: ")
    add_bullet("Límite de monto a crédito y límite máximo de garrafones vacíos prestados que se le permiten tener.", "• Condiciones Comerciales: ")

    add_h2("2.3. Rutas y Vehículos (/routes y /vehicles)")
    add_bullet("Crear los sectores territoriales de entrega (ej. Ruta Norte, Ruta Centro, Ruta Foránea).", "• Rutas: ")
    add_bullet("Registrar placas, marca, capacidad máxima de carga en garrafones y asignar el chofer/vendedor responsable habitual.", "• Vehículos: ")

    add_h2("2.4. Listas de Precios (/pricing)")
    add_bullet("Definición de precios diferenciados: Precio Público, Precio Distribuidor, Precio Especial Institucional.", "• Políticas de Precio: ")
    add_bullet("Asignación de listas de precios específicas por cliente o por zona.", "• Asignación: ")

    add_h1("3. Operación de Planta, Cargas y Liquidaciones")

    add_h2("3.1. Inventario Central (/inventory)")
    add_bullet("Registra el agua purificada recién embotellada lista para venta.", "• Entradas de Producción: ")
    add_bullet("Registra el ingreso de garrafones nuevos comprados a proveedores.", "• Adquisición de Envases: ")
    add_bullet("Muestra el stock disponible de producto terminado y envases vacíos listos para lavado.", "• Existencias Reales: ")

    add_h2("3.2. Asignación y Despacho de Cargas (/loads)")
    add_p("Procedimiento para dar salida al camión en la mañana:")
    add_bullet("Presiona el botón 'Nueva Carga'.", "Paso 1: ")
    add_bullet("Selecciona la Ruta, el Vehículo y el Vendedor responsable.", "Paso 2: ")
    add_bullet("Digita las cantidades de cada producto que suben físicamente al camión.", "Paso 3: ")
    add_bullet("Aprueba la carga. En ese instante se descuenta del inventario de bodega y se traslada al inventario móvil del camión.", "Paso 4: ")

    add_h2("3.3. Revisión y Aprobación de Liquidaciones (/settlements)")
    add_p("Procedimiento de cierre de tarde:")
    add_bullet("Revisar la liquidación enviada por el vendedor.", "Paso 1: ")
    add_bullet("Receptar físicamente el dinero en efectivo y cotejarlo contra el valor del sistema.", "Paso 2: ")
    add_bullet("Revisar el producto retornado sin vender y los envases vacíos devueltos a bodega.", "Paso 3: ")
    add_bullet("Aprobar la liquidación. Esto da por cerrada la carga y reingresa el stock sobrante a bodega central.", "Paso 4: ")

    add_h1("4. Cobranza, Créditos y Control de Garrafones")

    add_h2("4.1. Verificación de Transferencias Bancarias (/transfers)")
    add_bullet("Los vendedores cargan el número de comprobante o fotografía de la transferencia bancaria realizada por el cliente.", "• Recepción: ")
    add_bullet("El administrador valida contra la cuenta bancaria de la empresa.", "• Validación: ")
    add_bullet("Al presionar 'Aprobar', la venta queda formalmente saldada; si la transferencia no aparece en el banco, se rechaza y vuelve a saldo pendiente.", "• Resolución: ")

    add_h2("4.2. Créditos y Cuentas por Cobrar (/credit)")
    add_bullet("Reporte de clientes morosos o con cuentas por pagar.", "• Antigüedad de Saldos: ")
    add_bullet("Registro de cobros recibidos directamente en oficina central o en banco.", "• Abonos: ")

    add_h2("4.3. Control Estricto de Garrafones (/jugs)")
    add_p("Módulo para el control de activos retornables:")
    add_bullet("Muestra exactamente cuántos garrafones propiedad de la purificadora tiene cada cliente en su poder.", "• Saldo por Cliente: ")
    add_bullet("Permite identificar clientes con retención anormal de envases para exigir la devolución en la próxima visita.", "• Alertas de Pérdida: ")

    add_h1("5. Seguridad, Usuarios y Auditoría")

    add_h2("5.1. Anulaciones y Autorizaciones (/annulments y /operations-control)")
    add_bullet("Si un vendedor comete un error en un ticket, debe solicitar anulación con justificación.", "• Solicitud: ")
    add_bullet("Únicamente el ADMINISTRADOR tiene el privilegio de aprobar o denegar la anulación para evitar fraudes.", "• Aprobación: ")

    add_h2("5.2. Gestión de Usuarios (/administration)")
    add_bullet("Creación de nuevos empleados en el sistema.", "• Nuevos Usuarios: ")
    add_bullet("Asignación estricta del rol (ADMINISTRADOR, VENDEDOR, SUPERVISOR, BODEGA).", "• Roles: ")
    add_bullet("Restablecimiento de claves de acceso en caso de olvido.", "• Contraseñas: ")

    add_h2("5.3. Bitácora de Auditoría Inmutable (/audit)")
    add_bullet("Registra de forma inalterable quién hizo qué, a qué hora, desde qué dirección IP y qué valores fueron modificados.", "• Trazabilidad: ")
    add_bullet("Herramienta fundamental para revisiones fiscales, extravíos de producto o sospechas de alteración de datos.", "• Utilidad: ")

    add_h2("5.4. Datos de la Empresa (/company)")
    add_bullet("Configurar el Nombre Comercial, Logotipo corporativo, Moneda oficial (GTQ), Teléfono y Dirección que aparecen en todos los tickets y reportes.", "• Identidad Corporativa: ")

    doc.add_page_break()

    # =========================================================================
    # ANEXO: MATRIZ DE PERMISOS Y SEGURIDAD
    # =========================================================================
    sec3 = doc.add_paragraph()
    sec3.paragraph_format.space_before = Pt(10)
    sec3.paragraph_format.space_after = Pt(8)
    r3 = sec3.add_run("ANEXO: MATRIZ DE RESPONSABILIDADES Y ACCESO")
    r3.font.size = Pt(18)
    r3.font.bold = True
    r3.font.color.rgb = COLOR_PRIMARY

    add_p("Resumen comparativo de lo que puede y no puede realizar cada rol en el sistema:")

    table_p = doc.add_table(rows=12, cols=3)
    table_p.alignment = WD_TABLE_ALIGNMENT.CENTER
    heads_p = ["Módulo / Función", "Rol Vendedor", "Rol Administrador"]

    for i, h in enumerate(heads_p):
        cell = table_p.cell(0, i)
        set_cell_background(cell, "0F294A")
        set_cell_margins(cell, 100, 100, 120, 120)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    matriz = [
        ("Ver Carga de Ruta Asignada", "SÍ (Solo la suya)", "SÍ (Todas las rutas)"),
        ("Crear y Aprobar Nuevas Cargas", "NO", "SÍ"),
        ("Mi Recorrido y Mapa de Ruta", "SÍ", "SÍ (Modo supervisión)"),
        ("Registrar Ventas y Cobros", "SÍ (Stock de camión)", "SÍ (Cualquier cliente)"),
        ("Dar de Alta Nuevos Clientes", "Solo clientes ocasionales", "SÍ (Padrón completo con GPS)"),
        ("Modificar Precios de Productos", "NO", "SÍ"),
        ("Registrar Mermas de Ruta", "SÍ (Justificadas)", "SÍ"),
        ("Validar Transferencias Bancarias", "NO (Solo registrar boleta)", "SÍ (Aprobación final en banco)"),
        ("Aprobar Anulación de Tickets", "NO (Solo solicitar)", "SÍ"),
        ("Aprobar Liquidación Diaria", "NO (Solo enviar)", "SÍ (Cierre final y recepción de caja)"),
        ("Crear Usuarios y Asignar Roles", "NO", "SÍ")
    ]

    for row_idx, row_data in enumerate(matriz, start=1):
        bg = "F4F7FA" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, val in enumerate(row_data):
            cell = table_p.cell(row_idx, col_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, 80, 80, 120, 120)
            p = cell.paragraphs[0]
            p.add_run(val)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    add_callout("El cumplimiento de esta separación de funciones asegura la transparencia del negocio, evita descuadres en el inventario de agua y previene pérdidas de garrafones.", "CONCLUSIÓN OPERATIVA:")

    output_path = r"d:\UMG\Purificadora\Manual_de_Usuario_Purificadora.docx"
    doc.save(output_path)
    print(f"Manual generado con exito en: {output_path}")

if __name__ == '__main__':
    create_manual()
